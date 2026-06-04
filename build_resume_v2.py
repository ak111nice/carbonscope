# -*- coding: utf-8 -*-
"""生成第二版简历 —— 针对碳管理 SaaS 技术岗"""
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree

# ================================================================
# 格式化工具
# ================================================================
def _setup_cn(doc):
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Microsoft YaHei"
    font.size = Pt(10.5)
    rPr = style.element.get_or_add_rPr()
    rFonts = etree.SubElement(rPr, qn("w:rFonts"))
    rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

def _empty(p, text, bold=False, size=10.5):
    p = p if p else doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    return p

def _h(doc, text):
    h = doc.add_heading(text, level=2)
    return h

def _b(doc, text):
    doc.add_paragraph(text, style="List Bullet")

def _center(doc, text, size=10, color=None):
    p = doc.add_paragraph()
    p.alignment = 1
    r = p.add_run(text)
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor(*color)
    return p

# ================================================================
# 主内容
# ================================================================
doc = Document()
_setup_cn(doc)

# ── 头部 ──
p0 = doc.add_paragraph()
p0.alignment = 1
r0 = p0.add_run("周远林")
r0.font.size = Pt(18)
r0.bold = True

_center(doc, "电子科学与技术 本科 | 2026 届应届生 | 电话/邮箱请补充", color=(100, 100, 100))
_center(doc, "求职方向：碳数据工程师 / 碳管理平台开发 / SaaS 后端开发", size=9, color=(80, 80, 80))

# ── 个人优势 ──
_h(doc, "个人优势")

_b(doc, "扎实的碳管理知识：掌握 GB/T 32150 碳排放核算、ISO 14067 碳足迹、碳配额盈亏等核心业务逻辑")
_b(doc, "强技术落地能力：能用 Python 把碳核算流程自动化，独立开发完整产品并用 PyInstaller 打包成 exe")
_b(doc, "底层采集 + 上层开发兼顾：既懂 MODBUS 工业仪表，也能用 Chart.js/pdf.js 做数据可视化")
_b(doc, "从 0 到 1 闭环经验：72 小时竞赛独立完成方案 → 调试 → 联调全流程，同样适用于当前项目")

# ── 碳管理项目 ──
_h(doc, "核心项目")

_pj1 = doc.add_paragraph()
_empty(_pj1, "CarbonScope — AI 碳盘查自动化系统（独立全栈）｜ GB/T 32150 · ISO 14067", bold=True)
_empty(_pj1, "\n")
_empty(_pj1, "用 Python + JavaScript 搭建企业碳盘查系统，覆盖碳排放核算、配额盈亏模拟、产品碳足迹 LCA、PDF 账单解析、IoT 仪表采集、报告导出六大功能。", size=9.5)
_empty(_pj1, "\n")
_empty(_pj1, "PDF 自动提取：基于 pdf.js + 正则匹配中文关键词，从电费/燃气账单中精准提取用电量/用气量；带人工确认环节防止误读", size=9.5)
_empty(_pj1, "\n")
_empty(_pj1, "IoT 数据自动采集：按照 MODBUS 工业协议设计 7 通道仪表模拟器，自动读取并核算 → 碳排放实时可视化", size=9.5)
_empty(_pj1, "\n")
_empty(_pj1, "多端交付：纯前端 Web 版（index.html）+ PyInstaller 打包桌面 exe，零依赖双击即用", size=9.5)
_empty(_pj1, "\n")
_empty(_pj1, "GitHub 仓库: github.com/ak111nice/carbonscope ｜ 在线演示: ak111nice.github.io/carbonscope", size=9.5)

# ── 嵌入式项目 ──
_pj2 = doc.add_paragraph()
_empty(_pj2, "全国大学生电子设计竞赛（TI 杯）— 自行瞄准装置 | MSPM0G3507 + OpenMV + PID", bold=True)
_empty(_pj2, "\n")
_empty(_pj2, "72 小时独立完成方案设计到系统联调；负责硬件电路、传感器校准与 PID 参数整定", size=9.5)

# ── 专业技能 ──
_h(doc, "专业技能")

_b(doc, "碳核算：GB/T 32150 企业排放报告、ISO 14067 产品碳足迹、CBAM/欧盟碳关税基础")
_b(doc, "开发：Python（pandas/numpy/pdfplumber/Streamlit/PyInstaller）、JavaScript（Chart.js/pdf.js）")
_b(doc, "嵌入式/IoT：STM32/MSPM0、Keil/C、SPI/I2C/UART/CAN/MODBUS")
_b(doc, "工具链：Copilot、Claude Code、Git、Office")

# ── 教育背景 ──
_h(doc, "教育背景")
_p = doc.add_paragraph()
_empty(_p, "三江学院  电子科学与技术  本科 | 2022.09 - 2026.06")
_empty(_p, "\n")
_empty(_p, "主修：C语言、数据结构、模拟/数字电路、单片机、嵌入式系统设计、PCB设计", size=9.5)

# ── 证书 ──
_h(doc, "证书与语言")
_b(doc, "CET-4（英语四级）｜ 全国计算机等级考试二级 ｜ 普通话二级乙等")

# 保存
out = r"C:\Users\周远林\Desktop\碳中和项目\docs\周远林-碳管理技术岗简历-v2.docx"
doc.save(out)
print("Done: " + out)
