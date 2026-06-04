# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from lxml import etree

doc = Document()
style = doc.styles['Normal']
font = style.font
font.name = 'Microsoft YaHei'
font.size = Pt(11)
rPr = style.element.get_or_add_rPr()
rFonts = etree.SubElement(rPr, qn('w:rFonts'))
rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

def H(text, level=2):
    doc.add_heading(text, level=level)

def P(text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    return p

def B(text):
    doc.add_paragraph(text, style='List Bullet')

def T(headers, rows):
    t = doc.add_table(rows=len(rows)+1, cols=len(headers), style='Light Shading Accent 1')
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            t.rows[r+1].cells[c].text = cell
    return t

# ====== Title ======
H('碳中和 x AI 领域 21 天入行计划', 1)
P('目标：从零到可投递简历 + 有一个完整项目展示')
P('适用：电子科学与技术背景 + 会用 AI 工具(Codex/Claude Code)')

# ====== Week 1 ======
H('Week 1：碳核算基础（每天 2-3 小时）', 1)

P('Day 1-2：读懂核心公式', bold=True)
B('通读 GB/T 32150-2025 原文第4/5/6章(约20页): carbonnt.com/standards')
B('通读知乎《超全碳核算教程》: zhuanlan.zhihu.com/p/1909261733597516006')
B('背下核心公式: 碳排放 = 活动数据 x 排放因子 x GWP')
B('理解 Scope 1 / Scope 2 / Scope 3 的区别')

P('Day 3-4：看一份真实的碳报告', bold=True)
B('搜"上市公司 ESG 报告", 挑一份看(推荐: 宁德时代/宝钢/隆基)')
B('找到报告中"温室气体排放"章节, 看数据格式')
B('对比我们 demo 输出 vs 真实报告, 找差距')

P('Day 5：理解碳市场', bold=True)
B('了解全国碳市场基本原理: 配额制/CCER/履约周期')
B('了解 CBAM(欧盟碳边境调节机制)对中国出口企业影响')
B('搜索"生态环境部 碳市场扩容 2027"看最新政策')

# ====== Week 2 ======
H('Week 2：升级碳盘查项目（每天 3 小时）', 1)

P('Day 6-7：让系统能读取 PDF', bold=True)
B('用 pdfplumber 解析 PDF 格式的能源账单')
B('自动提取用电量/用气量等关键字段')
B('Bonus: 接一个 OCR(如果有扫描件需求)')

P('Day 8：加产品碳足迹计算', bold=True)
B('理解 LCA(全生命周期碳足迹)概念')
B('新增模块: 输入原材料用量/运输距离/用电量 -> 计算产品碳足迹')
B('参考标准: ISO 14067')

P('Day 9-10：加 AI 智能分析', bold=True)
B('用 LLM 对排放数据生成分析建议')
B('自动识别异常排放(某月突然升高 -> 自动标注)')
B('生成自然语言的减排建议(而非固定模板)')

P('Day 11：做前端展示(可选但加分)', bold=True)
B('用 Streamlit 做一个 Web 界面')
B('或至少用 matplotlib 做一个排放趋势图')
B('仪表盘: Scope 1/2/3 饼图/月度趋势/部门对比')

P('Day 12：代码整理 + README', bold=True)
B('完善 docstring 和注释')
B('写 README: 系统架构图/使用方法/技术栈')
B('录一个 30 秒 demo 视频(手机录屏即可)')

# ====== Week 3 ======
H('Week 3：包装 + 投递（每天 2 小时）', 1)

P('Day 13：简历改造', bold=True)
B('技能栏加: 碳核算(GB/T 32150)/LCA碳足迹/碳市场CBAM/Python')
B('项目经验: 写碳盘查系统 + GitHub 链接')

P('项目描述模板(直接抄到简历):')
P('独立开发基于 GB/T 32150-2025 的企业碳盘查自动化系统, 支持 PDF 账单解析、多源排放因子自动匹配、Scope 1/2 分项核算、AI 驱动的排放异常检测及合规性判断, 生成符合国标的碳核算报告。')

P('Day 14-15：搭建作品展示', bold=True)
B('项目推上 GitHub(README 写清楚 + demo 图或视频链接)')
B('写技术博客(掘金/知乎/CSDN): 标题《我用AI搭建了一个自动碳盘查系统》')
B('可选: LinkedIn 发英文版摘要')

P('Day 16-17：第一轮投递', bold=True)
B('BOSS直聘搜: "碳数据 开发" "能源 IoT" "碳 SaaS" "ESG 数据"')
B('猎聘搜: "碳管理 工程师" 筛选技术岗')

P('目标公司清单:', bold=True)
T(['类别', '公司'], [
    ['碳管理 SaaS', '碳阻迹/盟浪/碳衡科技/绿石'],
    ['能源 IoT', '阿里能耗宝/华为能源管理/远景能源'],
    ['新能源巨头数字化', '宁德时代/比亚迪/隆基'],
    ['央企信息化岗', '中石油昆仑数智/中海油能源发展'],
])

P('')
P('Day 18：第二波投递 + 内推', bold=True)
B('在脉脉/领英搜索目标公司员工, 私信说明你的背景+项目')
B('加入"碳圈青年""零碳未来"微信群(搜公众号)')

P('话术模板:')
P('我是一名电子工程应届生, 刚用AI工具做了一个自动碳盘查系统(GitHub: xxx), 对贵公司的碳管理产品方向很感兴趣, 能否交流一下?')

P('Day 19-20：面试准备', bold=True)
B('准备 3 分钟项目讲解(STAR 法则)')
B('准备: "为什么电子工程做碳中和?" -> 碳管理正从人工统计走向IoT自动采集+AI分析, 我的电子+AI背景恰好填补这个交叉缺口')
B('准备反问: "贵公司目前碳核算的数据来源是人工填报还是自动采集?"')

P('Day 21：复盘', bold=True)
B('记录投了多少份/收到多少面试')
B('根据面试反馈调整简历/项目')
B('循环: 持续改进项目 + 持续投递')

# ====== Resources ======
H('学习资源清单', 1)

T(['类别', '资源', '获取方式'], [
    ['核心标准', 'GB/T 32150-2025', 'carbonnt.com/standards'],
    ['核算教程', '知乎超全碳核算教程', 'zhuanlan.zhihu.com/p/1909261733597516006'],
    ['排放因子', '国家发改委省级清单指南', '搜索 PDF 下载'],
    ['碳市场', '上海环境能源交易所', 'cneeex.com'],
    ['国际标准', 'GHG Protocol 官网', 'ghgprotocol.org'],
    ['政策追踪', '生态环境部气候司', 'mee.gov.cn'],
    ['行业媒体', '36氪碳中和/碳阻迹公众号', '搜索关注'],
])

# ====== Checklist ======
H('每天检查清单', 1)
B('今天学的最重要的一个概念是什么?')
B('代码写了多少行? 有没有 commit?')
B('投了几份简历?')

out = r'C:\Users\周远林\Desktop\21天计划.docx'
doc.save(out)
print('Done: ' + out)
